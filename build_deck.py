"""Build 'Momentum Across Regimes — call with Gadi 29.5' deck.

Outputs momentum_across_regimes.pptx — import to Google Slides via
Drive upload → Open with Google Slides.
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

ROOT = Path("/Users/giladgang/momentum_regime")
OUT = ROOT / "momentum_across_regimes.pptx"

RED = RGBColor(0xCC, 0x00, 0x00)
GREEN = RGBColor(0x1B, 0x7A, 0x1B)
BLACK = RGBColor(0x10, 0x10, 0x10)
GREY = RGBColor(0x60, 0x60, 0x60)
LIGHT_GREY = RGBColor(0xB0, 0xB0, 0xB0)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_title(slide, text, top=Inches(0.5), left=Inches(0.6), width=Inches(12), size=36):
    tb = slide.shapes.add_textbox(left, top, width, Inches(1))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.name = "Arial"
    r.font.color.rgb = BLACK
    return tb


def add_bullets(slide, items, top=Inches(1.8), left=Inches(0.8), width=Inches(12), size=24, line_gap=18):
    tb = slide.shapes.add_textbox(left, top, width, Inches(5))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (text, color) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(line_gap)
        r = p.add_run()
        r.text = f"•  {text}"
        r.font.size = Pt(size)
        r.font.name = "Arial"
        r.font.color.rgb = color
    return tb


def add_text(slide, text, top, left, width, height, size=22, color=BLACK, align=PP_ALIGN.LEFT, bold=False, italic=False):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.name = "Arial"
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.italic = italic
    return tb


# ---------- Slide 1: Title ----------
s = prs.slides.add_slide(BLANK)
add_text(s, "Momentum Across Regimes", Inches(2.7), Inches(0.6), Inches(12), Inches(1.5), size=60, align=PP_ALIGN.CENTER)
add_text(s, "A regime-conditioned cross-sectional momentum model", Inches(3.9), Inches(0.6), Inches(12), Inches(0.6), size=22, color=GREY, align=PP_ALIGN.CENTER, italic=True)
add_text(s, "call with Gadi 29.5", Inches(4.7), Inches(0.6), Inches(12), Inches(0.6), size=22, color=GREY, align=PP_ALIGN.CENTER)


# ---------- Slides 2-5: Innovation build ----------
BULLETS = [
    ("legacy top 10% momentum, adjusting horizon weight or exposure according to market conditions", RED),
    ("ability to select any momentum level — maybe best to select the ones that crashed the hardest?", GREEN),
    ("simple probability indicator (market drawdown level)", RED),
    ("smart regime probability indicator (based on HMM)", GREEN),
]
for k in range(1, 5):
    s = prs.slides.add_slide(BLANK)
    add_title(s, "innovation in methodology — model flexibility and accuracy", size=30)
    visible = []
    for i, (txt, col) in enumerate(BULLETS):
        if i < k:
            visible.append((txt, col))
        else:
            visible.append((txt, RGBColor(0xFF, 0xFF, 0xFF)))
    add_bullets(s, visible, top=Inches(2.0), size=22, line_gap=24)


# ---------- Slide 6: Methodology ----------
s = prs.slides.add_slide(BLANK)
add_title(s, "methodology", size=36)
add_text(
    s,
    "HMM regime classifier + 12-month return horizons per stock → XGBoost ensemble → long–short portfolio",
    Inches(1.5), Inches(0.6), Inches(12), Inches(0.8), size=18, color=GREY,
)
hmm_path = ROOT / "plots" / "thesis" / "markov_chain_diagram.png"
if hmm_path.exists():
    s.shapes.add_picture(str(hmm_path), Inches(3.0), Inches(2.3), height=Inches(4.9))


# ---------- Slide 7: Main results — performance ----------
s = prs.slides.add_slide(BLANK)
add_title(s, "main results — performance", size=32)
add_text(
    s,
    "XGB (mom + π) Sharpe 1.11   |   FF6 α 24.7% (t = 4.78)   |   final $15.7 from $1 (2011–2025)",
    Inches(1.3), Inches(0.6), Inches(12), Inches(0.6), size=18, color=GREY,
)
perf_path = ROOT / "plots" / "thesis" / "cs_performance_regime_shaded.png"
if perf_path.exists():
    s.shapes.add_picture(str(perf_path), Inches(1.8), Inches(2.1), width=Inches(9.7))


# ---------- Slide 8 (new): Regime conditioning is the alpha ----------
s = prs.slides.add_slide(BLANK)
add_title(s, "regime conditioning is the alpha", size=32)
add_text(
    s,
    "the model earns its keep where legacy momentum dies",
    Inches(1.4), Inches(0.6), Inches(12), Inches(0.6), size=18, color=GREY, italic=True,
)
table_data = [
    ["Sharpe", "Full", "Calm (n=108)", "Panic (n=59)"],
    ["Market", "0.84", "0.64", "1.13"],
    ["Fixed 12-mo momentum", "0.06", "0.08", "0.05"],
    ["LR (linear baseline)", "-0.03", "-0.30", "0.35"],
    ["XGB (mom + π)", "1.11", "0.82", "1.56"],
]
rows, cols = len(table_data), 4
tbl_shape = s.shapes.add_table(rows, cols, Inches(2.2), Inches(2.4), Inches(9.0), Inches(3.2))
tbl = tbl_shape.table
col_widths = [Inches(3.0), Inches(2.0), Inches(2.0), Inches(2.0)]
for i, w in enumerate(col_widths):
    tbl.columns[i].width = w
for ri, row in enumerate(table_data):
    for ci, cell_text in enumerate(row):
        cell = tbl.cell(ri, ci)
        cell.text = ""
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = cell_text
        r.font.name = "Arial"
        r.font.size = Pt(18)
        is_header = ri == 0
        is_xgb = ri == 4
        r.font.bold = is_header or is_xgb
        r.font.color.rgb = BLACK if not is_header else GREY
add_text(
    s, "Panic Sharpe (1.56) ≈ 2× Calm Sharpe (0.82). The conditioning is doing the work.",
    Inches(5.9), Inches(0.7), Inches(12), Inches(0.5), size=16, color=GREEN, italic=True,
)


# ---------- Slide 9: Main results — buy-the-dip ----------
s = prs.slides.add_slide(BLANK)
add_title(s, "main results", size=36)
items = [
    ("in calm, invest in the stock with good momentum in most horizons", BLACK),
    ("in panic, invest in stocks that have underperformed (buying the dip)", BLACK),
    ("the risk here is sustained crash if there is no short rebound", BLACK),
]
add_bullets(s, items, top=Inches(1.6), size=22, line_gap=20)

add_text(s, "panic decomposes:", Inches(4.3), Inches(0.8), Inches(12), Inches(0.5), size=18, color=GREY, italic=True)
pan_rows = [
    ["Sub-type", "Months", "Ann Ret", "Sharpe"],
    ["Calm", "108", "13.8%", "0.82"],
    ["Panic — Crash (mkt down)", "18", "-7.2%", "-0.33"],
    ["Panic — Recovery (mkt up)", "41", "+64.4%", "2.35"],
]
tbl_shape = s.shapes.add_table(len(pan_rows), 4, Inches(0.9), Inches(4.9), Inches(8.5), Inches(2.0))
tbl = tbl_shape.table
widths = [Inches(3.4), Inches(1.4), Inches(1.7), Inches(1.4)]
for i, w in enumerate(widths):
    tbl.columns[i].width = w
for ri, row in enumerate(pan_rows):
    for ci, cell_text in enumerate(row):
        cell = tbl.cell(ri, ci)
        cell.text = ""
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = cell_text
        r.font.name = "Arial"
        r.font.size = Pt(14)
        r.font.bold = ri == 0
        r.font.color.rgb = GREY if ri == 0 else BLACK
add_text(
    s, "panic Sharpe is driven entirely by recovery months", Inches(7.0), Inches(0.9),
    Inches(12), Inches(0.5), size=14, color=GREY, italic=True,
)


# ---------- Slide 10 (new): Open questions ----------
s = prs.slides.add_slide(BLANK)
add_title(s, "open questions for Gadi", size=32)
qs = [
    ("how to hedge the sustained-crash tail — short rebound never arrives?", BLACK),
    ("which momentum level to pick in panic — the worst crashers, or something risk-adjusted?", BLACK),
    ("is the HMM regime probability stable enough to trade on at month boundaries?", BLACK),
    ("international replication (UK / JP panels) — does the regime signal port?", BLACK),
]
add_bullets(s, qs, top=Inches(1.9), size=22, line_gap=22)


prs.save(OUT)
print(f"wrote {OUT}")
