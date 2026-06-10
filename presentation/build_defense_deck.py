"""Build the thesis defense deck -> presentation/momentum_defense.pptx."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.text import PP_ALIGN

from presentation import content as C
from presentation import deck_lib as L
from presentation.rasterize_assets import ensure_pngs

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / C.FIG
OUT = ROOT / "presentation" / "momentum_defense.pptx"


def _fig(name):
    return FIG / name


# ---------------- Main slides ----------------

def s01_title(prs, A):
    s = L.new_slide(prs)
    L.figure_fit(s, A["logo"], Inches(4.9), Inches(0.7), Inches(3.5), Inches(1.6))
    tf = L._box(s, Inches(0.8), Inches(2.7), Inches(11.7), Inches(2.0)).text_frame
    L._set(tf.paragraphs[0], C.TITLE, 40, color=L.ACCENT, bold=True, align=PP_ALIGN.CENTER)
    L._set(tf.add_paragraph(), C.SUBTITLE, 22, color=L.MUTE, italic=True, align=PP_ALIGN.CENTER)
    tf2 = L._box(s, Inches(0.8), Inches(4.7), Inches(11.7), Inches(2.2)).text_frame
    for line in (C.AUTHOR, C.SUPERVISOR, C.PROGRAMME, C.DATE):
        p = tf2.paragraphs[0] if not tf2.paragraphs[0].runs else tf2.add_paragraph()
        L._set(p, line, 16, color=L.INK, align=PP_ALIGN.CENTER)
    L.notes(s, "Title and framing. Aim: understand momentum across regimes, not beat the market.")


def s02_motivation(prs, A):
    s = L.new_slide(prs)
    L.title(s, "Momentum: most profitable anomaly, most fragile")
    L.bullets(s, [
        "Buying past winners and shorting past losers is one of finance's most robust anomalies.",
        ("Yet it crashes hard at regime turns - losses over 50% in the 2009 rebound.", L.BAD),
    ], top=Inches(1.5), width=Inches(12))
    L.stat(s, C.MDD_FIX12, "fixed 12-mo momentum\nmax drawdown", Inches(1.6), Inches(3.6), color=L.BAD)
    L.stat(s, "0.06", "its Sharpe\n(2011-2024)", Inches(5.5), Inches(3.6), color=L.BAD)
    L.stat(s, "?", "what selects stocks\nin each regime", Inches(9.2), Inches(3.6), color=L.ACCENT)
    L.takeaway(s, "Question: how does the cross-section reorganize between calm and panic, and what selects stocks in each?", color=L.ACCENT)
    L.notes(s, "Set up the puzzle. Existing fixes scale exposure or blend horizons; the selection rule stays unconditional. We ask a different question.")


def s03_literature(prs, A):
    s = L.new_slide(prs)
    L.title(s, "Prior work, and the gap")
    cols = [
        ("Momentum & crashes", ["Jegadeesh-Titman (1993)", "Daniel-Moskowitz (2016):\ncrashes are predictable"]),
        ("Crash management", ["Adjusts EXPOSURE or HORIZON", "Barroso-Santa-Clara; GHM (2023)"]),
        ("Regime + ML", ["Cooper (2004): coarse up/down", "Gu (2020); Beckmeyer-\nWiedemann (2025)"]),
    ]
    x = Inches(0.7)
    for head, items in cols:
        tf = L._box(s, x, Inches(1.6), Inches(3.9), Inches(3.2)).text_frame
        L._set(tf.paragraphs[0], head, 18, color=L.ACCENT, bold=True)
        for it in items:
            L._set(tf.add_paragraph(), it, 14, color=L.INK)
        x += Inches(4.05)
    L.takeaway(s, "Gap: nobody studies how the full momentum term structure reorganizes across regimes, or whether capturing it needs nonlinear regime-conditioned selection.", color=L.ACCENT)
    L.notes(s, "Three buckets. Each treats the regime as a switch on sizing or horizon, never as a feature inside a nonlinear selection rule. That gap is the thesis.")


def s04_methodology_hmm(prs, A):
    s = L.new_slide(prs)
    L.title(s, "Methodology I - the regime signal (Stage 1)")
    L.pipeline(s, ["HMM\n(4 stress features)", "pi : panic\nprobability", "XGBoost", "Long-short\nportfolio"], top=Inches(1.35))
    L.figure_fit(s, _fig("regime_probabilities.png"), Inches(1.2), Inches(2.6), Inches(11.0), Inches(3.9))
    L.takeaway(s, "A 2-state Bayesian HMM turns four stress indicators into a continuous panic probability pi that updates monthly. pi is a context variable, not a return predictor.")
    L.notes(s, "pi spikes at dot-com, GFC, COVID, 2022. Gibbs/FFBS detail is in backup B1. Trained 1990-2010, frozen for 2011-2024.")


def s05_methodology_xgb(prs, A):
    s = L.new_slide(prs)
    L.title(s, "Methodology II - cross-sectional selection (Stage 2)")
    L.bullets(s, [
        "Inputs: 12 momentum horizons (1-12 months) + the regime probability pi.",
        ("XGBoost ranks stocks on PREDICTED forward return, not raw momentum.", L.ACCENT),
        "So the long leg need not hold high-momentum names; pi routes how the whole term structure is used.",
        "Benchmarked against a deterministic formula (DET) and a logistic regression (LR).",
    ], top=Inches(1.6), width=Inches(12))
    L.table(s, [["Method", "Sharpe"],
                ["DET (formula)", f"{C.SHARPE_DET:.2f}"],
                ["LR (linear)", f"{C.SHARPE_LR:.2f}"],
                ["XGB (nonlinear)", f"{C.SHARPE_XGB:.2f}"]],
            Inches(0.8), Inches(4.9), Inches(4.6), Inches(1.7), highlight_row=3)
    L.takeaway(s, "Only the nonlinear model turns these inputs into performance. XGB specifics (depth-4, 50-seed ensemble) are in backup B12.")
    L.notes(s, "Emphasise predicted-return ranking - that is what lets the long leg flip composition by regime. DET and LR share the exact same inputs and collapse.")


def s06_results(prs, A):
    s = L.new_slide(prs)
    L.title(s, "Headline results (2011-2024, out-of-sample)")
    L.figure_fit(s, _fig("cs_performance_regime_shaded.png"), Inches(0.7), Inches(1.5), Inches(7.6), Inches(5.0))
    L.table(s, [["Sharpe", ""],
                ["Market", f"{C.SHARPE_MARKET:.2f}"],
                ["Fixed 12-mo mom.", f"{C.SHARPE_FIX12:.2f}"],
                ["XGB (mom + pi)", f"{C.SHARPE_XGB:.2f}"]],
            Inches(8.7), Inches(1.7), Inches(3.9), Inches(1.9), highlight_row=3)
    tf = L._box(s, Inches(8.7), Inches(4.0), Inches(4.2), Inches(2.4)).text_frame
    L._set(tf.paragraphs[0], f"Six-factor alpha {C.FF6_ALPHA}", 20, color=L.GOOD, bold=True)
    L._set(tf.add_paragraph(), f"t = {C.FF6_T}", 16, color=L.MUTE)
    L._set(tf.add_paragraph(), f"Remove pi -> Sharpe {C.SHARPE_NO_PI:.2f}", 14, color=L.INK)
    L._set(tf.add_paragraph(), f"Fixed mom. MDD {C.MDD_FIX12}", 14, color=L.BAD)
    L.takeaway(s, "XGB is the only row clearing every conventional t-stat threshold; the linear baseline on identical features collapses.")
    L.notes(s, "Credibility first. Alpha survives FF6 (cancels mechanical winner tilt). The advantage concentrates in the shaded panic months - that motivates the mechanism.")


def s07_mechanism_flip(prs, A):
    s = L.new_slide(prs)
    L.title(s, "The mechanism I - the direction flip")
    L.figure_fit(s, _fig("zscore_long_heatmap.png"), Inches(3.3), Inches(1.5), Inches(6.7), Inches(5.0))
    tf = L._box(s, Inches(0.6), Inches(2.0), Inches(2.6), Inches(4.0)).text_frame
    L._set(tf.paragraphs[0], "Calm", 20, color=L.GOOD, bold=True)
    L._set(tf.add_paragraph(), "picks sit ABOVE the cross-sectional mean (winners)", 14, color=L.INK)
    L._set(tf.add_paragraph(), "Panic", 20, color=L.BAD, bold=True)
    L._set(tf.add_paragraph(), "picks sit BELOW at every horizon (the beaten-down)", 14, color=L.INK)
    L.takeaway(s, "'Buy the dip' emerges only after conditioning on pi - no single-horizon momentum rule would select these names.")
    L.notes(s, "Each row is a test month; columns are the 12 horizons. Red below the mean in panic, blue above in calm. Foreshadow the four modes.")


def s08_mechanism_clusters(prs, A):
    s = L.new_slide(prs)
    L.title(s, "The mechanism II - four modes, one market cycle")
    panels = ["zscore_l2_k4_panel_c0.png", "zscore_l2_k4_panel_c1.png",
              "zscore_l2_k4_panel_c2.png", "zscore_l2_k4_panel_c3.png"]
    coords = [(Inches(0.7), Inches(1.6)), (Inches(6.9), Inches(1.6)),
              (Inches(0.7), Inches(4.0)), (Inches(6.9), Inches(4.0))]
    for name, (lx, ty), cname, csh in zip(panels, coords, C.CLUSTER_NAMES, C.CLUSTER_SHARPES):
        L.figure_fit(s, _fig(name), lx, ty, Inches(5.6), Inches(2.2))
        cap = L._box(s, lx, ty - Inches(0.05), Inches(5.6), Inches(0.3)).text_frame
        L._set(cap.paragraphs[0], f"{cname}  (Sharpe {csh:.2f})", 12, color=L.ACCENT, bold=True)
    L.takeaway(s, "Calm continuation -> transition -> post-panic recovery -> deep crisis. Momentum where it works, reversal where it crashes.")
    L.notes(s, "The heatmap clusters into four selection modes that trace one cycle. Cluster 4 (deep crisis) earns the highest Sharpe, 1.82, by buying the most beaten-down names.")


def s09_novel(prs, A):
    s = L.new_slide(prs)
    L.title(s, "Why it's novel")
    L.bullets(s, [
        (f"pi is a routing variable: {C.SHAP_PI} of total SHAP, ~6x its 1/13 share.", L.ACCENT),
        f"Nonlinearity is the binding constraint: DET {C.SHARPE_DET:.2f}, LR {C.SHARPE_LR:.2f}, XGB {C.SHARPE_XGB:.2f}.",
        "Composition, not exposure: which stocks are held flips by regime, not how much.",
    ], top=Inches(1.6), width=Inches(12))
    L.table(s, [["CAPM beta", "Calm", "Panic"],
                ["Long leg", f"{C.LEG_BETA_LONG[0]:.2f}", f"{C.LEG_BETA_LONG[1]:.2f}"],
                ["Short leg", f"{C.LEG_BETA_SHORT[0]:.2f}", f"{C.LEG_BETA_SHORT[1]:.2f}"]],
            Inches(0.8), Inches(4.4), Inches(6.0), Inches(1.8))
    L.takeaway(s, "Long-leg beta exceeds short-leg beta in every regime - the strategy rides rebounds instead of being crushed by them, unlike D&M/Barroso exposure scaling.")
    L.notes(s, "This is the contribution slide. Contrast each literature bucket: we change WHICH stocks, not HOW MUCH, and we learn it nonlinearly from a continuous real-time signal.")


def s10_robustness(prs, A):
    s = L.new_slide(prs)
    L.title(s, "Robustness")
    L.bullets(s, [
        (f"Survives the 2009 momentum crash: XGB {C.CRASH2009_XGB} vs unconditional {C.CRASH2009_FIX}.", L.GOOD),
        "Alpha survives all five factor models (CAPM through FF6).",
    ], top=Inches(1.5), width=Inches(12))
    L.table(s, [["Sharpe (2011-2024)", "XGB", "Fixed 12-mo"],
                ["US", f"{C.SHARPE_XGB:.2f}", f"{C.SHARPE_FIX12:.2f}"],
                ["UK", f"{C.UK_XGB:.2f}", f"{C.UK_FIX:.2f}"],
                ["Japan", f"{C.JP_XGB:.2f}", f"{C.JP_FIX:.2f}"]],
            Inches(0.8), Inches(3.6), Inches(6.4), Inches(2.3), highlight_row=0)
    tf = L._box(s, Inches(7.6), Inches(3.6), Inches(5.2), Inches(2.4)).text_frame
    L._set(tf.paragraphs[0], "Max drawdown compresses:", 16, color=L.INK, bold=True)
    L._set(tf.add_paragraph(), f"US {C.MDD_US}  |  UK {C.MDD_UK}  |  JP {C.MDD_JP}", 15, color=L.GOOD)
    L._set(tf.add_paragraph(), "vs fixed-momentum -62% to -72%", 14, color=L.BAD)
    L.takeaway(s, "Works where momentum already works (US, UK) and where it does not (Japan, baseline ~0). Only the HMM features are region-specific.")
    L.notes(s, "Japan is the strongest evidence: baseline momentum is absent there, yet regime-conditioning produces 0.57. Rules out 'just a better momentum implementation'.")


def s11_conclusions(prs, A):
    s = L.new_slide(prs)
    L.title(s, "Conclusions & future work")
    cols = [
        ("Contributions", ["Novel filter + nonlinear model combination",
                            "Concrete stock-level mechanism (4 modes)"]),
        ("Limitation", ["Reversal-failure in a sustained bear",
                        f"dot-com: {C.DOTCOM_XGB} ({C.DOTCOM_MDD} MDD)"]),
        ("Future work", ["Regime forecasting / trajectory",
                         "Neural nets; D&M exposure overlay"]),
    ]
    x = Inches(0.7)
    for head, items in cols:
        tf = L._box(s, x, Inches(1.7), Inches(3.9), Inches(3.5)).text_frame
        L._set(tf.paragraphs[0], head, 18, color=L.ACCENT, bold=True)
        for it in items:
            L._set(tf.add_paragraph(), it, 15, color=L.INK)
        x += Inches(4.05)
    L.takeaway(s, "Momentum's cross-section reorganizes across regimes; a continuous regime signal routes a nonlinear model between winner selection and reversal.")
    L.notes(s, "Close on the two contributions and one future direction. Be honest about the dot-com reversal-failure mode and the D&M circuit-breaker complement.")


MAIN_BUILDERS = [s01_title, s02_motivation, s03_literature, s04_methodology_hmm,
                 s05_methodology_xgb, s06_results, s07_mechanism_flip,
                 s08_mechanism_clusters, s09_novel, s10_robustness, s11_conclusions]


def _backup_builder(entry):
    title_text, figure, bullet_items, notes_text = entry

    def build(prs, A):
        fig_path = None
        if figure and figure.startswith("ASSET:"):
            fig_path = A[figure.split(":", 1)[1]]
        elif figure:
            fig_path = FIG / figure
        L.content_slide(
            prs, title_text=title_text, figure_path=fig_path,
            bullets_items=bullet_items, notes_text=notes_text,
        )
    return build


BACKUP_BUILDERS = [_backup_builder(e) for e in C.BACKUP_SLIDES]


def _divider(prs, text):
    s = L.new_slide(prs)
    tf = L._box(s, Inches(0.8), Inches(3.0), Inches(11.7), Inches(1.5)).text_frame
    L._set(tf.paragraphs[0], text, 32, color=L.ACCENT, bold=True, align=PP_ALIGN.CENTER)


def build_presentation():
    prs = Presentation()
    prs.slide_width, prs.slide_height = L.SLIDE_W, L.SLIDE_H
    A = ensure_pngs()
    for build in MAIN_BUILDERS:
        build(prs, A)
    _divider(prs, "Backup slides  -  for the Q&A")
    for build in BACKUP_BUILDERS:
        build(prs, A)
    return prs


def main():
    prs = build_presentation()
    prs.save(str(OUT))
    print(f"wrote {OUT}  ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
