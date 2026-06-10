from presentation import content as C


def test_canonical_headline_numbers():
    assert C.SHARPE_XGB == 1.11
    assert C.FF6_ALPHA == "24.1%"
    assert C.FF6_T == "4.81"
    assert C.PANIC_SHARPE == 1.53
    assert C.CALM_SHARPE == 0.84
    assert C.SHARPE_FIX12 == 0.06
    assert C.MDD_FIX12 == "-72.2%"
    assert C.CLUSTER_SHARPES == [0.92, 0.93, 1.21, 1.82]


def test_stale_gadi_numbers_are_not_used():
    # The old "call with Gadi" deck used a different cut; these must never appear.
    forbidden = ["24.7%", "1.56", "$15.7"]
    blob = " ".join(str(v) for v in vars(C).values() if isinstance(v, (str, int, float, list)))
    for bad in forbidden:
        assert bad not in blob, f"stale value {bad!r} present in content.py"


def test_primitives_build_a_titled_slide_with_notes():
    from pptx import Presentation
    from pptx.util import Inches
    from presentation import deck_lib as L
    prs = Presentation()
    prs.slide_width, prs.slide_height = L.SLIDE_W, L.SLIDE_H
    s = L.content_slide(
        prs, title_text="Hello",
        bullets_items=["one", ("two", L.GOOD)],
        table_rows=[["A", "B"], ["1", "2"]],
        takeaway_text="done", notes_text="say hi",
    )
    titles = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
    assert "Hello" in titles
    assert s.notes_slide.notes_text_frame.text == "say hi"
    assert len(prs.slides) == 1


def test_rasterize_produces_pngs():
    from presentation.rasterize_assets import ensure_pngs
    paths = ensure_pngs()
    assert set(paths) >= {"logo", "gibbs", "acf", "riskav"}
    for key, p in paths.items():
        assert p.exists() and p.stat().st_size > 1000, f"{key} png missing/empty"
