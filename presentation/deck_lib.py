"""Reusable python-pptx primitives for the defense deck (16:9, clean-minimal)."""
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
FONT = "Calibri"

INK    = RGBColor(0x1A, 0x1A, 0x1A)
MUTE   = RGBColor(0x6B, 0x6B, 0x6B)
ACCENT = RGBColor(0x00, 0x2C, 0x5F)   # Tilburg navy
GOOD   = RGBColor(0x1B, 0x7A, 0x3D)
BAD    = RGBColor(0xB3, 0x2A, 0x2A)
PANEL  = RGBColor(0xEE, 0xF1, 0xF6)


def new_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _box(slide, left, top, width, height):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tb.text_frame.word_wrap = True
    return tb


def _set(p, text, size, color=INK, bold=False, italic=False, align=PP_ALIGN.LEFT):
    p.alignment = align
    r = p.add_run()
    r.text = text
    f = r.font
    f.size, f.name, f.color.rgb, f.bold, f.italic = Pt(size), FONT, color, bold, italic
    return r


def title(slide, text, color=ACCENT):
    tb = _box(slide, Inches(0.6), Inches(0.4), Inches(12.1), Inches(1.0))
    _set(tb.text_frame.paragraphs[0], text, 30, color=color, bold=True)
    return tb


def takeaway(slide, text, color=GOOD):
    tb = _box(slide, Inches(0.6), Inches(6.75), Inches(12.1), Inches(0.5))
    _set(tb.text_frame.paragraphs[0], text, 16, color=color, italic=True)
    return tb


def bullets(slide, items, left=Inches(0.7), top=Inches(1.7), width=Inches(12), size=20, gap=12):
    tf = _box(slide, left, top, width, Inches(5)).text_frame
    for i, item in enumerate(items):
        text, color = item if isinstance(item, tuple) else (item, INK)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        _set(p, "•  " + text, size, color=color)


def figure_fit(slide, path, box_left, box_top, box_w, box_h):
    """Place an image inside a box, preserving aspect ratio, centered."""
    iw, ih = Image.open(str(path)).size
    if iw / ih > box_w / box_h:
        w = box_w; h = int(box_w * ih / iw)
    else:
        h = box_h; w = int(box_h * iw / ih)
    slide.shapes.add_picture(str(path), box_left + (box_w - w) // 2,
                             box_top + (box_h - h) // 2, width=w, height=h)


def table(slide, rows, left, top, width, height, highlight_row=None, font=14):
    nrow, ncol = len(rows), len(rows[0])
    gfx = slide.shapes.add_table(nrow, ncol, left, top, width, height).table
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = gfx.cell(ri, ci)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            align = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.RIGHT
            head = ri == 0
            hi = highlight_row is not None and ri == highlight_row
            _set(p, str(val), font, color=MUTE if head else INK, bold=head or hi, align=align)
    return gfx


def pipeline(slide, stages, top=Inches(1.5), height=Inches(0.85),
             left=Inches(0.9), total_w=Inches(11.5)):
    n = len(stages)
    gap = Inches(0.35)
    box_w = (total_w - gap * (n - 1)) // n
    x = left
    for i, label in enumerate(stages):
        sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, top, box_w, height)
        sh.fill.solid(); sh.fill.fore_color.rgb = PANEL
        sh.line.color.rgb = ACCENT; sh.line.width = Pt(1.25)
        sh.text_frame.word_wrap = True
        _set(sh.text_frame.paragraphs[0], label, 13, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
        if i < n - 1:
            ar = _box(slide, x + box_w, top, gap, height)
            _set(ar.text_frame.paragraphs[0], "→", 22, color=MUTE, align=PP_ALIGN.CENTER)
        x += box_w + gap


def stat(slide, number, label, left, top, width=Inches(3.6), color=ACCENT):
    tf = _box(slide, left, top, width, Inches(1.9)).text_frame
    _set(tf.paragraphs[0], number, 44, color=color, bold=True, align=PP_ALIGN.CENTER)
    _set(tf.add_paragraph(), label, 14, color=MUTE, align=PP_ALIGN.CENTER)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def content_slide(prs, *, title_text, figure_path=None, bullets_items=None,
                  table_rows=None, table_highlight=None, takeaway_text=None, notes_text=None):
    """Generic single-figure / bullets / table slide used by most backup slides."""
    s = new_slide(prs)
    title(s, title_text)
    if bullets_items:
        bullets(s, bullets_items, top=Inches(1.6))
    if figure_path:
        if bullets_items:
            figure_fit(s, figure_path, Inches(6.9), Inches(1.7), Inches(6.0), Inches(4.6))
        else:
            figure_fit(s, figure_path, Inches(1.2), Inches(1.6), Inches(11.0), Inches(4.9))
    if table_rows:
        table(s, table_rows, Inches(0.8), Inches(4.8), Inches(7.6), Inches(1.8),
              highlight_row=table_highlight)
    if takeaway_text:
        takeaway(s, takeaway_text)
    if notes_text:
        notes(s, notes_text)
    return s
